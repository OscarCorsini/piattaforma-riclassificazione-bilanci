# -*- coding: utf-8 -*-
"""
excel_writer.py
===============
Scrittura dei dati riclassificati nel file Excel "Template" tramite
openpyxl, senza alterare formule o formattazioni preesistenti.

Il modulo rileva DINAMICAMENTE la struttura del template (righe delle
categorie, colonne degli anni) invece di usare un mapping statico, cosi'
funziona con template reali che possono differire da quello di esempio:

  1. Trova la riga di ciascuna categoria cercando l'etichetta testuale
     nel foglio (es. "Corrispettivi normali").
  2. Rileva la colonna corrispondente a ciascun anno leggendo i valori
     calcolati da Excel nella riga di intestazione degli anni (utile
     quando l'anno e' il risultato di una formula, es. "anno base + 1").
  3. Se la cella di destinazione fa parte di un intervallo di celle unite
     (merged), scrive nella cella "ancora" (in alto a sinistra).
  4. Se una categoria di ENTRATA o USCITA non viene trovata nel template,
     prova a INSERIRE dinamicamente una nuova riga per ospitarla (vedi
     aggiungi_categoria_dinamica), cosi' l'utente puo' scrivere una voce
     completamente nuova nel wizard di revisione e vederla comparire come
     riga a se stante nel file Excel finale. Se anche l'inserimento
     dinamico fallisce (o si tratta di gestione fiscale, che vive in
     un'area a celle fisse e non viene mai toccata dinamicamente), la
     categoria viene segnalata come "non mappata" invece di generare un
     errore bloccante.

Supporta inoltre la scrittura di PIU' ANNI in un'unica esecuzione
(popola_template_multi), aprendo il workbook una sola volta e scrivendo
tutti gli anni richiesti prima di salvare.
"""

from __future__ import annotations

import io
import re
import unicodedata
from copy import copy
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.cell_range import CellRange

from config import SHEET_NAME_RICLASSIFICAZIONE, TEMPLATE_PATH
from groq_client import RiclassificazioneResult


class ExcelTemplateError(Exception):
    """Errore nell'apertura o nella struttura del file Excel template."""


class ExcelMappingError(Exception):
    """Errore nel mapping tra l'anno selezionato e le colonne del template."""


@dataclass
class WriteReport:
    """Riepilogo di cio' che e' stato scritto, utile per un feedback chiaro in UI."""
    celle_scritte: int
    voci_non_mappate: List[str]
    foglio: str


# ---------------------------------------------------------------------------
# UTILITA' DI LETTURA / RICERCA NEL FOGLIO
# ---------------------------------------------------------------------------
def _get_sheet(workbook, sheet_name: str):
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    return workbook.active


def _normalizza_etichetta(valore) -> str:
    if valore is None:
        return ""
    testo = str(valore)
    testo = unicodedata.normalize("NFKD", testo)
    testo = testo.encode("ascii", "ignore").decode("ascii")
    testo = re.sub(r"[^a-zA-Z0-9]+", " ", testo)
    return re.sub(r"\s+", " ", testo).strip().lower()


def _trova_riga_categoria(sheet, categoria: str, max_row: int = 300, max_col: int = 6) -> Optional[int]:
    target = _normalizza_etichetta(categoria)
    if not target:
        return None
    limite_righe = min(sheet.max_row or max_row, max_row)
    limite_colonne = min(sheet.max_column or max_col, max_col)
    for row in range(1, limite_righe + 1):
        for col in range(1, limite_colonne + 1):
            valore = sheet.cell(row=row, column=col).value
            if valore is not None and _normalizza_etichetta(valore) == target:
                return row
    return None


def _scansiona_riga_anni(sheet_valori, row: int, limite_colonne: int) -> Dict[str, str]:
    """Legge una singola riga e restituisce {colonna_lettera: anno} per i
    valori che sembrano plausibili anni (interi 1900-2100)."""
    colonne_riga: Dict[str, str] = {}
    for col in range(1, limite_colonne + 1):
        valore = sheet_valori.cell(row=row, column=col).value
        if isinstance(valore, bool):
            continue
        if isinstance(valore, (int, float)):
            anno_int = int(valore)
            if 1900 <= anno_int <= 2100 and float(anno_int) == float(valore):
                colonne_riga[get_column_letter(col)] = str(anno_int)
    return colonne_riga


def _rileva_colonne_anno(sheet_valori, max_row: int = 30, max_col: int = 40) -> Dict[str, str]:
    """Restituisce le colonne-anno {colonna_lettera: anno}.

    Normalmente basta la PRIMA riga (dall'alto) con almeno 2 valori interi
    verosimili come anno. Alcuni template pero' hanno in quella riga valori
    "vecchi"/non aggiornati (es. due colonne diverse con lo stesso anno,
    perche' il foglio e' stato copiato/modificato senza ricalcolare i valori
    statici): in quel caso la riga contiene un anno duplicato su piu'
    colonne, segno che non e' affidabile al 100%. Quando succede, si cerca
    nelle righe immediatamente successive (di solito la riga con le formule
    di calcolo dell'anno, es. "=C3+1") una riga SENZA duplicati che possa
    correggere le colonne in conflitto, e si integrano i risultati.
    """
    limite_righe = min(sheet_valori.max_row or max_row, max_row)
    limite_colonne = min(sheet_valori.max_column or max_col, max_col)

    for row in range(1, limite_righe + 1):
        colonne_riga = _scansiona_riga_anni(sheet_valori, row, limite_colonne)
        if len(colonne_riga) < 2:
            continue

        valori = list(colonne_riga.values())
        duplicati = {v for v in valori if valori.count(v) > 1}
        if not duplicati:
            return colonne_riga

        # La riga trovata ha valori duplicati (es. due colonne con lo
        # stesso anno): prova a correggere le colonne in conflitto usando
        # le prossime righe, che spesso contengono la sequenza corretta
        # calcolata da formule (es. anno_base+1, +2, ...). La correzione e'
        # iterativa: sistemare una colonna puo' rivelare un nuovo conflitto
        # (es. la colonna successiva aveva "ereditato" lo stesso anno), quindi
        # si ripete finche' non ci sono piu' duplicati o non si puo' piu'
        # correggere nulla con la riga corrente.
        for succ_row in range(row + 1, min(row + 6, limite_righe + 1)):
            riga_succ = _scansiona_riga_anni(sheet_valori, succ_row, limite_colonne)
            if len(riga_succ) < 2:
                continue
            valori_succ = list(riga_succ.values())

            while True:
                valori_attuali = list(colonne_riga.values())
                duplicati_attuali = {v for v in valori_attuali if valori_attuali.count(v) > 1}
                if not duplicati_attuali:
                    break
                corretto = False
                for col, val in list(colonne_riga.items()):
                    if val in duplicati_attuali and col in riga_succ:
                        nuovo_valore = riga_succ[col]
                        if nuovo_valore != val and valori_succ.count(nuovo_valore) == 1:
                            colonne_riga[col] = nuovo_valore
                            corretto = True
                if not corretto:
                    break

            # Integra anche eventuali colonne-anno aggiuntive non presenti
            # nella riga originale (es. anni futuri/previsionali), purche'
            # non creino un nuovo duplicato.
            for col, valore in riga_succ.items():
                if col not in colonne_riga and valore not in colonne_riga.values():
                    colonne_riga[col] = valore

        return colonne_riga

    return {}


def _trova_cella_ancora(sheet, cella_unita):
    for rng in sheet.merged_cells.ranges:
        if cella_unita.coordinate in rng:
            return sheet.cell(row=rng.min_row, column=rng.min_col)
    return None


def _leggi_bytes(file_like) -> bytes:
    try:
        file_like.seek(0)
    except Exception:
        pass
    dati = file_like.read()
    try:
        file_like.seek(0)
    except Exception:
        pass
    return dati


# ---------------------------------------------------------------------------
# INSERIMENTO DINAMICO DI UNA NUOVA RIGA CATEGORIA (voci non a mappatura fissa)
# ---------------------------------------------------------------------------
_CELL_REF_RE = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")


def _shift_refs_ge(valore, soglia: int, delta: int):
    """Nelle formule (stringhe che iniziano con '='), incrementa di
    'delta' ogni riferimento di riga >= soglia. Lascia invariato tutto il
    resto (valori non-formula, riferimenti a righe precedenti la soglia)."""
    if not isinstance(valore, str) or not valore.startswith("="):
        return valore

    def repl(match):
        col_abs, col, row_abs, row_s = match.groups()
        row_num = int(row_s)
        if row_num >= soglia:
            row_num += delta
        return f"{col_abs}{col}{row_abs}{row_num}"

    return _CELL_REF_RE.sub(repl, valore)


def _retarget_refs(valore, riga_da, riga_a):
    """Nelle formule, sostituisce i riferimenti che puntano ESATTAMENTE a
    riga_da con riga_a (usato per copiare la formula di una riga modello
    in una nuova riga, mantenendo pero' intatti i riferimenti alle altre
    righe, es. ai totali)."""
    if not isinstance(valore, str) or not valore.startswith("="):
        return valore

    def repl(match):
        col_abs, col, row_abs, row_s = match.groups()
        row_num = int(row_s)
        if row_num == riga_da:
            row_num = riga_a
        return f"{col_abs}{col}{row_abs}{row_num}"

    return _CELL_REF_RE.sub(repl, valore)


def _inserisci_riga_manuale(sheet, insert_at: int, riga_modello: int, nuova_label_col: int, nuova_label: str) -> int:
    """
    Inserisce una nuova riga in posizione insert_at (tutto cio' che si
    trova a partire da insert_at scende di una posizione), copiando lo
    stile e le formule della riga_modello (che si trova immediatamente
    sopra insert_at, e quindi NON viene spostata) nella nuova riga.

    Usa una tecnica a snapshot manuale invece di sheet.insert_rows():
    quest'ultima puo' perdere silenziosamente il contenuto di celle unite
    complesse (bug noto di openpyxl), come gia' verificato e risolto in
    precedenza per la correzione strutturale del template.

    Restituisce il numero della nuova riga (== insert_at).
    """
    max_row = sheet.max_row or 1
    max_col = sheet.max_column or 1

    # 1) Snapshot di ogni cella non vuota (valore + stile).
    snapshot = {}
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            cell = sheet.cell(row=row, column=col)
            if isinstance(cell, MergedCell):
                continue
            if cell.value is None:
                continue
            snapshot[(row, col)] = {
                "value": cell.value,
                "font": copy(cell.font),
                "fill": copy(cell.fill),
                "border": copy(cell.border),
                "alignment": copy(cell.alignment),
                "number_format": cell.number_format,
            }

    # 2) Snapshot dei merge e delle altezze riga.
    merged_ranges = [str(rng) for rng in sheet.merged_cells.ranges]
    row_heights = {row: sheet.row_dimensions[row].height for row in range(1, max_row + 1)}

    # 3) Rimuove tutti i merge e svuota tutte le celle.
    for rng_str in list(merged_ranges):
        sheet.unmerge_cells(rng_str)
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            cell = sheet.cell(row=row, column=col)
            if not isinstance(cell, MergedCell):
                cell.value = None

    # 4) Riscrive ogni cella nella nuova posizione, shiftando le formule.
    for (row, col), dati in snapshot.items():
        nuova_riga = row + 1 if row >= insert_at else row
        valore = _shift_refs_ge(dati["value"], insert_at, 1)
        cell = sheet.cell(row=nuova_riga, column=col)
        cell.value = valore
        cell.font = dati["font"]
        cell.fill = dati["fill"]
        cell.border = dati["border"]
        cell.alignment = dati["alignment"]
        cell.number_format = dati["number_format"]

    # 5) Ricrea i merge nelle posizioni corrette.
    for rng_str in merged_ranges:
        rng = CellRange(rng_str)
        min_row = rng.min_row + 1 if rng.min_row >= insert_at else rng.min_row
        max_row_r = rng.max_row + 1 if rng.max_row >= insert_at else rng.max_row
        if min_row == max_row_r and rng.min_col == rng.max_col:
            continue
        sheet.merge_cells(start_row=min_row, start_column=rng.min_col, end_row=max_row_r, end_column=rng.max_col)

    # 6) Ripristina le altezze delle righe.
    for row, height in row_heights.items():
        nuova_riga = row + 1 if row >= insert_at else row
        if height is not None:
            sheet.row_dimensions[nuova_riga].height = height
    altezza_modello = sheet.row_dimensions[riga_modello].height
    if altezza_modello is not None:
        sheet.row_dimensions[insert_at].height = altezza_modello

    # 7) Copia il contenuto della riga_modello (ancora al suo posto,
    # perche' riga_modello < insert_at) nella nuova riga vuota, ri-
    # targettizzando solo i riferimenti che puntavano esattamente alla
    # riga_modello (es. "=D29/$D$24" diventa "=D30/$D$24" sulla riga
    # nuova, mantenendo pero' $D$24 intatto).
    for col in range(1, max_col + 1):
        cell_modello = sheet.cell(row=riga_modello, column=col)
        if cell_modello.value is None:
            continue
        valore = _retarget_refs(cell_modello.value, riga_modello, insert_at)
        cell_nuova = sheet.cell(row=insert_at, column=col)
        cell_nuova.value = valore
        cell_nuova.font = copy(cell_modello.font)
        cell_nuova.fill = copy(cell_modello.fill)
        cell_nuova.border = copy(cell_modello.border)
        cell_nuova.alignment = copy(cell_modello.alignment)
        cell_nuova.number_format = cell_modello.number_format

    # Replica anche gli eventuali merge di cui faceva parte la riga_modello
    # (es. "D29:E29") sulla nuova riga.
    for rng_str in merged_ranges:
        rng = CellRange(rng_str)
        if rng.min_row == riga_modello and rng.max_row == riga_modello and rng.min_col != rng.max_col:
            sheet.merge_cells(start_row=insert_at, start_column=rng.min_col, end_row=insert_at, end_column=rng.max_col)

    # Etichetta della nuova voce (sovrascrive il testo copiato dal modello).
    sheet.cell(row=insert_at, column=nuova_label_col).value = nuova_label

    return insert_at


def _estendi_somma_categoria(sheet, riga_sum: int, vecchia_fine: int, nuova_fine: int) -> None:
    """Estende l'intervallo delle formule SUM (o simili) presenti sulla
    riga_sum, spostando la fine dell'intervallo da vecchia_fine a
    nuova_fine (es. "=SUM(D29:E39)" diventa "=SUM(D29:E40)"), cosi' la
    riga appena inserita viene inclusa nel totale."""
    pattern = re.compile(rf":([A-Z]{{1,3}}){vecchia_fine}\b")
    max_col = sheet.max_column or 1
    for col in range(1, max_col + 1):
        cell = sheet.cell(row=riga_sum, column=col)
        valore = cell.value
        if isinstance(valore, str) and valore.startswith("=") and pattern.search(valore):
            cell.value = pattern.sub(rf":\g<1>{nuova_fine}", valore)


def aggiungi_categoria_dinamica(sheet, sezione: str, nuova_categoria: str) -> Optional[int]:
    """
    Inserisce una nuova riga per una categoria di ENTRATA o USCITA non
    presente nel template (l'utente ha scritto una voce completamente
    nuova nel wizard di revisione, che non corrisponde a nessuna
    categoria esistente).

      - sezione "entrata": la riga viene inserita subito PRIMA di
        "Totale entrate", ed estende la somma di quella riga.
      - sezione "uscita": la riga viene inserita come nuovo "di cui"
        subito PRIMA di "Affitti" (nel blocco "Costi normali"), ed
        estende la somma di "Costi normali".

    La gestione fiscale non e' MAI gestita da questa funzione: vive in
    un'area a celle fisse separata del template e non deve mai essere
    alterata dinamicamente.

    Restituisce il numero della nuova riga se l'inserimento e' riuscito,
    altrimenti None (il chiamante deve trattare la categoria come "non
    mappata" senza che il file venga alterato).
    """
    if sezione == "entrata":
        riga_ancora = _trova_riga_categoria(sheet, "Totale entrate")
        etichetta_col = 2
        riga_sum = riga_ancora
    elif sezione == "uscita":
        riga_ancora = _trova_riga_categoria(sheet, "Affitti")
        etichetta_col = 3
        riga_sum = _trova_riga_categoria(sheet, "Costi normali")
    else:
        return None

    if riga_ancora is None or riga_sum is None:
        return None

    riga_modello = riga_ancora - 1
    if riga_modello < 1:
        return None
    if sezione == "uscita" and riga_modello <= riga_sum:
        # Per le uscite riga_sum ("Costi normali") e' una riga diversa da
        # riga_ancora ("Affitti"): se la riga modello finisse dentro o
        # prima del blocco somma, la struttura del template non e' quella
        # attesa. Per le entrate invece riga_sum COINCIDE con riga_ancora
        # ("Totale entrate" e' essa stessa la riga di somma), quindi questo
        # controllo non si applica.
        return None

    try:
        nuova_riga = _inserisci_riga_manuale(
            sheet,
            insert_at=riga_ancora,
            riga_modello=riga_modello,
            nuova_label_col=etichetta_col,
            nuova_label=nuova_categoria,
        )
        riga_sum_finale = riga_sum + 1 if riga_sum >= riga_ancora else riga_sum
        _estendi_somma_categoria(sheet, riga_sum_finale, riga_modello, riga_ancora)
    except Exception:
        return None

    return nuova_riga


def _scrivi_valori_anno(sheet, colonna: str, risultato: RiclassificazioneResult) -> WriteReport:
    """Scrive i valori di UN anno (una colonna) gia' individuata nel foglio
    gia' aperto. Funzione interna riusata sia dal percorso a singolo anno
    che da quello multi-anno."""
    celle_scritte = 0
    voci_non_mappate: List[str] = []

    def _scrivi_cella(riga: int, categoria: str) -> bool:
        cella = sheet[f"{colonna}{riga}"]
        nonlocal celle_scritte
        if isinstance(cella, MergedCell):
            ancora = _trova_cella_ancora(sheet, cella)
            if ancora is None:
                return False
            cella = ancora
        try:
            cella.value = valore
            celle_scritte += 1
            return True
        except Exception as exc:
            raise ExcelTemplateError(
                f"Errore scrivendo la categoria '{categoria}' nella cella '{cella.coordinate}': {exc}"
            ) from exc

    # Entrate e uscite: se la categoria non e' gia' presente nel template,
    # si prova a inserire dinamicamente una nuova riga per ospitarla.
    for sezione, valori_da_scrivere in (
        ("entrata", risultato.totale_entrate()),
        ("uscita", risultato.totale_uscite()),
    ):
        for categoria, valore in valori_da_scrivere.items():
            riga = _trova_riga_categoria(sheet, categoria)
            if riga is None:
                riga = aggiungi_categoria_dinamica(sheet, sezione, categoria)
            if riga is None or not _scrivi_cella(riga, categoria):
                voci_non_mappate.append(categoria)

    # Gestione fiscale: MAI inserita dinamicamente (area a celle fisse).
    for categoria, valore in risultato.totale_gestione_fiscale().items():
        riga = _trova_riga_categoria(sheet, categoria)
        if riga is None or not _scrivi_cella(riga, categoria):
            voci_non_mappate.append(categoria)

    return WriteReport(
        celle_scritte=celle_scritte,
        voci_non_mappate=voci_non_mappate,
        foglio=sheet.title,
    )


# ---------------------------------------------------------------------------
# FUNZIONE PUBBLICA: TEMPLATE PREDEFINITO
# ---------------------------------------------------------------------------
def carica_template_predefinito() -> bytes:
    """Legge dal disco il template Excel fisso distribuito con l'app."""
    try:
        with open(TEMPLATE_PATH, "rb") as f:
            return f.read()
    except Exception as exc:
        raise ExcelTemplateError(
            f"Impossibile leggere il template predefinito ({TEMPLATE_PATH}): {exc}"
        ) from exc


# ---------------------------------------------------------------------------
# NUOVO FOGLIO: DETTAGLIO MICROVOCI
# ---------------------------------------------------------------------------
def _scrivi_dettaglio_voci(workbook, risultati_per_anno: Dict[str, RiclassificazioneResult]):
    """Crea un nuovo foglio per riversare l'elenco di tutte le microvoci divise per anno."""
    sheet_name = "Dettaglio Microvoci"
    if sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        sheet.delete_rows(1, sheet.max_row)
    else:
        sheet = workbook.create_sheet(title=sheet_name)

    headers = ["Anno", "Sezione", "Macro Categoria", "Voce Originale", "Importo"]
    sheet.append(headers)

    from openpyxl.styles import Font
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for anno, risultato in risultati_per_anno.items():
        for macro, voci in risultato.entrate.items():
            for voce in voci:
                sheet.append([anno, "Entrate", macro, voce.descrizione, voce.importo])
        for macro, voci in risultato.uscite.items():
            for voce in voci:
                sheet.append([anno, "Uscite", macro, voce.descrizione, voce.importo])
        for macro, voci in risultato.gestione_fiscale.items():
            for voce in voci:
                sheet.append([anno, "Gestione Fiscale", macro, voce.descrizione, voce.importo])

    for col in sheet.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        adjusted_width = (max_length + 2)
        sheet.column_dimensions[column].width = adjusted_width


# ---------------------------------------------------------------------------
# FUNZIONE PUBBLICA: RILEVAMENTO ANNI DISPONIBILI
# ---------------------------------------------------------------------------
def rileva_anni_disponibili(template_bytes: bytes) -> List[str]:
    try:
        workbook_valori = openpyxl.load_workbook(io.BytesIO(template_bytes), data_only=True)
        sheet_valori = _get_sheet(workbook_valori, SHEET_NAME_RICLASSIFICAZIONE)
        colonne_anno = _rileva_colonne_anno(sheet_valori)
        return sorted(set(colonne_anno.values()))
    except Exception:
        return []


# ---------------------------------------------------------------------------
# FUNZIONE PUBBLICA: SCRITTURA MULTI-ANNO (usata dall'interfaccia)
# ---------------------------------------------------------------------------
def popola_template_multi(
    template_bytes: bytes,
    risultati_per_anno: Dict[str, RiclassificazioneResult],
) -> Tuple[io.BytesIO, Dict[str, WriteReport], Dict[str, str]]:
    """
    Apre il template UNA SOLA VOLTA e scrive i dati di piu' anni nello
    stesso file, restituendo un unico buffer scaricabile. Aggiunge anche un
    foglio "Dettaglio Microvoci" con ogni singola voce originale, utile per
    verificare la classificazione.

    Restituisce (buffer, report_per_anno, errori_per_anno): gli anni scritti
    con successo compaiono in report_per_anno, quelli per cui non e' stata
    trovata una colonna nel template compaiono in errori_per_anno (senza
    interrompere l'elaborazione degli altri anni).
    """
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(template_bytes), data_only=False, keep_links=True)
        workbook_valori = openpyxl.load_workbook(io.BytesIO(template_bytes), data_only=True)
    except InvalidFileException as exc:
        raise ExcelTemplateError(
            "Il file template non e' un file Excel (.xlsx) valido."
        ) from exc
    except Exception as exc:
        raise ExcelTemplateError(
            f"Impossibile aprire il file Excel template: {exc}"
        ) from exc

    # Forza Excel a ricalcolare TUTTE le formule all'apertura del file:
    # openpyxl non calcola le formule mentre scrive, quindi le celle di
    # subtotale/totale (es. "Acquisti ordinari", "Totale uscite", "M.O.L.")
    # manterrebbero il valore memorizzato nel template originale finche'
    # non viene forzato un ricalcolo completo, risultando vuote o non
    # aggiornate quando l'utente apre il file scaricato.
    try:
        workbook.calculation.fullCalcOnLoad = True
    except Exception:
        pass

    sheet = _get_sheet(workbook, SHEET_NAME_RICLASSIFICAZIONE)
    sheet_valori = _get_sheet(workbook_valori, SHEET_NAME_RICLASSIFICAZIONE)

    colonne_anno = _rileva_colonne_anno(sheet_valori)
    mappa_anno_colonna = {anno_val: col for col, anno_val in colonne_anno.items()}

    report_per_anno: Dict[str, WriteReport] = {}
    errori_per_anno: Dict[str, str] = {}

    for anno, risultato in risultati_per_anno.items():
        colonna = mappa_anno_colonna.get(str(anno))
        if colonna is None:
            anni_trovati = ", ".join(sorted(mappa_anno_colonna.keys())) or "nessuno"
            errori_per_anno[anno] = (
                f"Il template non contiene una colonna per l'anno '{anno}'. "
                f"Anni individuati automaticamente nel template: {anni_trovati}."
            )
            continue
        report_per_anno[anno] = _scrivi_valori_anno(sheet, colonna, risultato)

    _scrivi_dettaglio_voci(workbook, risultati_per_anno)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return buffer, report_per_anno, errori_per_anno


# ---------------------------------------------------------------------------
# FUNZIONE PUBBLICA: SCRITTURA A SINGOLO ANNO (compatibilita')
# ---------------------------------------------------------------------------
def popola_template_excel(template_file, anno: str, risultato: RiclassificazioneResult):
    """Variante a singolo anno, mantenuta per compatibilita': internamente
    usa lo stesso motore multi-anno con un solo elemento."""
    dati = _leggi_bytes(template_file)
    buffer, report_per_anno, errori_per_anno = popola_template_multi(dati, {anno: risultato})

    if anno in errori_per_anno:
        raise ExcelMappingError(errori_per_anno[anno])

    return buffer, report_per_anno[anno]
